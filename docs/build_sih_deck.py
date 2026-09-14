"""Rebuild slides 2-4 of the SIH pitch deck per DeepTrek team requirements.

Slide 2 — challenges on LEFT, solution + 2 flowcharts + 6 steps on RIGHT.
Slide 3 — compact 4-block tech stack + APIx formula on LEFT, HLD (cache + NN) on RIGHT.
Slide 4 — Phase 1 = 4 metro cities; Future scope with 2 images + concise chips.
"""
import sys
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

SRC = Path(r"C:\Users\LENOVO\Downloads\_Real_ Time_ Airfare.pptx")
OUT = Path(r"C:\Users\LENOVO\Downloads\Real_Time_Airfare_SIH_pitch_deck_DeepTrek.pptx")
IMG = Path(r"C:\Users\LENOVO\Desktop\SIH-2026\docs\deck-images")

# colours --------------------------------------------------------------
BLUE_D = RGBColor(0x0D, 0x47, 0xA1)
BLUE_M = RGBColor(0x15, 0x65, 0xC0)
BLUE_L = RGBColor(0xE3, 0xF2, 0xFD)
GREEN_D = RGBColor(0x1B, 0x5E, 0x20)
GREEN_M = RGBColor(0x2E, 0x7D, 0x32)
GREEN_L = RGBColor(0xE8, 0xF5, 0xE9)
ORANGE_D = RGBColor(0xE6, 0x51, 0x00)
ORANGE_L = RGBColor(0xFF, 0xF3, 0xE0)
PURPLE_D = RGBColor(0x6A, 0x1B, 0x9A)
PURPLE_L = RGBColor(0xF3, 0xE5, 0xF5)
TEAL_M = RGBColor(0x00, 0x69, 0x5C)
TEAL_L = RGBColor(0xE0, 0xF2, 0xF1)
GREY = RGBColor(0x5F, 0x63, 0x68)
INK = RGBColor(0x1A, 0x1A, 0x2E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Segoe UI"


# helpers ---------------------------------------------------------------
def _text(slide, x, y, w, h, parts, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True, sp_after=2):
    """parts = list of (list-of-runs, level). run = (text, size, bold, color)."""
    box = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    first = True
    for runs, level in parts:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.level = level
        p.space_after = Pt(sp_after)
        for text, size, bold, color in runs:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = FONT
            r.font.color.rgb = color
    return box


def _line(slide, x, y, w, h, text, size=13, color=INK, bold=True):
    return _text(slide, x, y, w, h, [([(text, size, bold, color)], 0)])


def _pill(slide, x, y, w, h, text, fill, tcolor=WHITE, size=16):
    r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    r.adjustments[0] = 0.5
    r.fill.solid(); r.fill.fore_color.rgb = fill
    r.line.color.rgb = fill; r.line.width = Pt(0.5)
    r.shadow.inherit = False
    tf = r.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Emu(20000)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run = p.add_run(); run.text = text
    run.font.size = Pt(size); run.font.bold = True; run.font.name = FONT; run.font.color.rgb = tcolor
    return r


def _card(slide, x, y, w, h, accent, title, desc, tsize=15, dsize=11.5, fill=WHITE, bar_w=90000):
    r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    r.adjustments[0] = 0.07
    r.fill.solid(); r.fill.fore_color.rgb = fill
    r.line.color.rgb = accent; r.line.width = Pt(1.1)
    r.shadow.inherit = False

    bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x + 40000), Emu(y + 60000),
                                 Emu(bar_w), Emu(h - 120000))
    bar.adjustments[0] = 0.5
    bar.fill.solid(); bar.fill.fore_color.rgb = accent
    bar.line.fill.background(); bar.shadow.inherit = False

    _line(slide, x + 320000, y + 70000, w - 380000, 200000, title, size=tsize, color=accent)
    _text(slide, x + 320000, y + 310000, w - 380000, h - 380000, [([(desc, dsize, False, INK)], 0)])
    return r


def _chip(slide, x, y, w, h, num, label, accent, size=11.5):
    r = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    r.adjustments[0] = 0.18
    r.fill.solid(); r.fill.fore_color.rgb = WHITE
    r.line.color.rgb = accent; r.line.width = Pt(1.0)
    r.shadow.inherit = False
    d = 110000
    cx, cy = x + 95000, y + (h - d) // 2
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(cx), Emu(cy), Emu(d), Emu(d))
    c.fill.solid(); c.fill.fore_color.rgb = accent
    c.line.fill.background(); c.shadow.inherit = False
    tfc = c.text_frame; tfc.vertical_anchor = MSO_ANCHOR.MIDDLE
    pc = tfc.paragraphs[0]; pc.alignment = PP_ALIGN.CENTER
    rc = pc.add_run(); rc.text = num
    rc.font.size = Pt(12); rc.font.bold = True; rc.font.name = FONT; rc.font.color.rgb = WHITE
    _text(slide, cx + d + 20000, y, w - (cx + d + 20000 - x) - 20000, h,
          [([(label, size, True, INK)], 0)], anchor=MSO_ANCHOR.MIDDLE)
    return r


def _pic(slide, img, x, y, w, h):
    return slide.shapes.add_picture(str(IMG / img), Emu(x), Emu(y), Emu(w), Emu(h))


prs = Presentation(str(SRC))
slides = list(prs.slides)


def delete_except(slide, keep):
    for shape in list(slide.shapes):
        if shape.name not in keep:
            shape._element.getparent().remove(shape._element)


# ═══════════════════════════════════════════════════════════════════════
# SLIDE 2  —  Problem (left) | Solution + flowcharts (right)
# ═══════════════════════════════════════════════════════════════════════
s2 = slides[1]
delete_except(s2, {"object 2", "object 5", "object 6", "object 7", "object 8", "object 11"})

# column divider (titles sit above y≈1.4M)
div = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(6096000), Emu(1450000), Emu(26000), Emu(4700000))
div.fill.solid(); div.fill.fore_color.rgb = RGBColor(0xB0, 0xBE, 0xC5)
div.line.fill.background(); div.shadow.inherit = False

# LEFT — challenges ------------------------------------------------
_pill(s2, 250000, 1500000, 5700000, 330000, "THE PROBLEM",
      BLUE_D, size=17)
_line(s2, 250000, 1920000, 5600000, 300000,
      "Airlines change fares dynamically; India has no official airfare price index.",
      size=12.5, color=INK, bold=False)

challs = [
    ("1", "Fragmented Sources", "Airfares are scattered across airlines, OTAs & GDS feeds — no single trustworthy source.",
     BLUE_M),
    ("2", "Dynamic Pricing", "Fares swing 200–400% within a single day; manual sampling can't track rapid movement.",
     ORANGE_D),
    ("3", "Non-Standardized Data", "Fare components, taxes & formats differ across every platform.",
     GREEN_M),
    ("4", "Limited High-Frequency Measurement", "The current CPI model can't capture route & lead-time price behaviour.",
     PURPLE_D),
]
pos = [(250000, 2350000), (3090000, 2350000), (250000, 3450000), (3090000, 3450000)]
for (num, title, desc, accent), (x, y) in zip(challs, pos):
    _card(s2, x, y, 2690000, 990000, accent, title, desc, tsize=15.5, dsize=11.5)

_line(s2, 250000, 4530000, 5600000, 300000,
      "Unlike CPI/WPI, no published index measures how the cost of domestic air travel moves.",
      size=12, color=GREY, bold=False)

# RIGHT — solution + flowcharts -------------------------------------
_pill(s2, 6250000, 1500000, 5700000, 330000, "OUR SOLUTION — AIRINDEX",
      GREEN_D, size=17)
_line(s2, 6250000, 1920000, 5700000, 300000,
      "A transparent, reproducible price index over a fixed basket of metro routes — for MoSPI / NSO / RBI.",
      size=12.5, color=INK, bold=False)

_pic(s2, "02-scraping.png", 6250000, 2350000, 2790000, 1569000)
_pic(s2, "03-cleaning.png", 9200000, 2350000, 2790000, 1569000)
_line(s2, 6250000, 3960000, 2790000, 200000, "1 · Automated collection — Playwright (Python) · rate-limited",
      size=10.5, color=BLUE_M, bold=False)
_line(s2, 9200000, 3960000, 2790000, 200000, "2 · Cleaning & validation — Pandas · IQR outliers",
      size=10.5, color=GREEN_M, bold=False)

steps = [
    ("01", "Automated Collection", BLUE_M),
    ("02", "Data Cleaning & Validation", GREEN_M),
    ("03", "Fare Normalization", ORANGE_D),
    ("04", "Statistical Index Engine", PURPLE_D),
    ("05", "Real-Time Airfare Price Index", TEAL_M),
    ("06", "Dashboard + API", BLUE_D),
]
xs = [6250000, 8160000, 10070000]
ys = [4300000, 4920000]
for i, (num, label, accent) in enumerate(steps):
    _chip(s2, xs[i % 3], ys[i // 3], 1810000, 470000, num, label, accent, size=11)

# ═══════════════════════════════════════════════════════════════════════
# SLIDE 3  —  Tech stack + APIx formula (left) | HLD architecture (right)
# ═══════════════════════════════════════════════════════════════════════
s3 = slides[2]
delete_except(s3, {"object 2", "object 5", "object 6", "object 7", "object 8", "object 36"})

div3 = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(6096000), Emu(1450000), Emu(26000), Emu(4700000))
div3.fill.solid(); div3.fill.fore_color.rgb = RGBColor(0xB0, 0xBE, 0xC5)
div3.line.fill.background(); div3.shadow.inherit = False

# LEFT — compact 4-block tech stack ---------------------------------
_pill(s3, 250000, 1500000, 5700000, 330000, "TECH STACK — 4 LAYERS", BLUE_D, size=17)

stack = [
    ("DATA COLLECTION", "Playwright (Python) · headless Chrome", "rate-limited 60 req/h · retries · APScheduler daily jobs", BLUE_M),
    ("DATA PROCESSING", "Pandas · NumPy · Python", "dedup · null-fare handling · IQR (1.5x) outlier flag · quality score", GREEN_M),
    ("FRONTEND", "Streamlit + Plotly", "composite trend · route analytics · heatmap · raw explorer", ORANGE_D),
    ("DATABASE", "SQLite (data/apix.db)", "raw_flights · cleaned_flights · daily / weekly / monthly index", PURPLE_D),
]
spos = [(250000, 1980000), (3090000, 1980000), (250000, 2900000), (3090000, 2900000)]
for (title, sub, desc, accent), (x, y) in zip(stack, spos):
    _card(s3, x, y, 2690000, 790000, accent, title, f"{sub}\n{desc}", tsize=14.5, dsize=10.5)

# APIx math strip
_pill(s3, 250000, 3830000, 1790000, 260000, "HOW APIx IS CALCULATED", ORANGE_D, size=11.5)
_pic(s3, "api-formula.png", 250000, 4150000, 5450000, 1299000)
_text(s3, 250000, 5520000, 5600000, 300000, [
    ([("Base day = 100 · Daily / weekly / monthly frequencies · DGCA + booking-lead weights · config-driven", 11, False, GREY)], 0),
])

# RIGHT — HLD architecture -------------------------------------------
_pill(s3, 6250000, 1500000, 5700000, 330000, "ARCHITECTURE — HIGH-LEVEL DESIGN (HLD)", GREEN_D, size=16)
_line(s3, 6250000, 1910000, 5700000, 200000,
      "Storage with cache layer · neural net on the consumer layer", size=12, color=INK, bold=False)
_pic(s3, "hld-architecture.png", 6250000, 2180000, 5700000, 3224000)

_chip(s3, 6250000, 5750000, 2740000, 360000, "\u2194", "CACHE LAYER — TTL 60 s · materialised views", GREEN_M, size=10)
_chip(s3, 9160000, 5750000, 2780000, 360000, "NN", "NN ON CONSUMER LAYER — TensorFlow fare classifier", PURPLE_D, size=10)

# ═══════════════════════════════════════════════════════════════════════
# SLIDE 4  —  Phase 1 = 4 cities · Future scope with 2 images + chips
# ═══════════════════════════════════════════════════════════════════════
s4 = slides[3]

# Phase 1 body -> 4 metro cities
for shape in s4.shapes:
    if shape.name == "Text 3":
        shape.text_frame.text = "4 Metro Cities · 12 Routes\n-> Automated data pipeline -> Fare normalization -> Airfare Index -> Interactive dashboard"
    if shape.name in ("Text 9", "Picture 38"):   # old future-scope bullets + right image
        shape._element.getparent().remove(shape._element)

# future scope: concise chips + 2 visuals (below the existing FUTURE SCOPE header at y=3603437)
futs = [
    ("AI/ML", "Anomaly detection + fare-direction forecasting", PURPLE_D),
    ("DATA", "More OTA sources · 50+ routes · DGCA back-test", ORANGE_D),
    ("SERVE", "FastAPI for NSO / RBI · auto-refresh scheduler", TEAL_M),
]
fx = [577032, 4500000, 8420000]
for (tag, desc, accent), x in zip(futs, fx):
    ch = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(3970000), Emu(3300000), Emu(440000))
    ch.adjustments[0] = 0.16
    ch.fill.solid(); ch.fill.fore_color.rgb = WHITE
    ch.line.color.rgb = accent; ch.line.width = Pt(1.0)
    ch.shadow.inherit = False
    _text(s4, x + 70000, 4040000, 3100000, 440000,
          [([(tag, 10, True, accent)], 0), ([(f"  {desc}", 11, False, INK)], 0)],
          anchor=MSO_ANCHOR.MIDDLE)

# two visuals: frontend dashboard mock + upcoming charts
_line(s4, 577032, 4455000, 2489000, 130000, "Frontend preview — Streamlit + Plotly dashboard",
      size=11, color=BLUE_M)
_pic(s4, "dashboard-mock.png", 577032, 4605000, 2489000, 1400000)

_line(s4, 3300000, 4455000, 2980000, 130000, "Upcoming analytics — charts, elasticity, route trends",
      size=11, color=ORANGE_D)
_pic(s4, "trend-charts.png", 3300000, 4605000, 2980000, 1400000)

prs.save(str(OUT))
print("SAVED:", OUT)
print("Slides:", len(prs.slides))